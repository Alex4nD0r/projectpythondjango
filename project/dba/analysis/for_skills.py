import csv
import itertools

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Border, Side
import matplotlib.pyplot as plt
import numpy as np
from openpyxl.worksheet.table import Table, TableStyleInfo


def get_skills(file_name):
    skills_list = []
    skills = {}
    sorted_skills = {}

    with open(file_name, mode='r', encoding='utf-8-sig')as file:
        reader = csv.reader(file)
        for row in reader:
            if row[0] in ['Администратор баз данных', 'баз данных', 'оператор баз данных', 'базы данных',
                          'oracle', 'mysql', 'data base', 'database', 'dba', 'bd', 'бд', 'базами данны']:
                if row[1] == '':
                    continue
                skills_list.append(row[1])

    skills_list = ''.join(skills_list)
    skills_list = skills_list.split('\n')

    for skill in skills_list:
        skills[skill] = skills_list.count(skill)

    sorted_values = sorted(skills.values(), reverse=True)
    for i in sorted_values:
        for k in skills.keys():
            if skills[k] == i:
                sorted_skills[k] = i
                break

    sorted_skills = dict(itertools.islice(sorted_skills.items(), 10))

    return sorted_skills


def generate_excel(file_name):
    skills_dict = get_skills(file_name)
    workbook = Workbook()
    worksheet = workbook.active
    data = []
    worksheet.append(['Навык', 'Количество'])

    for key, value in skills_dict.items():
        data.append([key, value])

    for row in data:
        worksheet.append(row)

    tab = Table(displayName="Skills", ref='A1:B10')

    style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                           showLastColumn=False, showRowStripes=True, showColumnStripes=True)
    tab.tableStyleInfo = style

    worksheet.add_table(tab)
    workbook.save("report_skills.xlsx")


def generate_image(file_name):
    skills_dict = get_skills(file_name)
    fig, (ax3) = plt.subplots(nrows=1, ncols=1)

    ax3.set_title('Навык', fontdict={'fontsize': 8})
    ax3.barh(list([str(a).replace(' ', '\n').replace('-', '-\n') for a in reversed(list(skills_dict.keys()))]),
             list(reversed(list(skills_dict.values()))), color='blue', height=0.5, align='center')
    ax3.yaxis.set_tick_params(labelsize=6)
    ax3.xaxis.set_tick_params(labelsize=8)
    ax3.grid(axis='x')

    plt.tight_layout()
    plt.savefig('graph_skills.png')


def main(file_name):
    generate_image(file_name)
    generate_excel(file_name)


main('vacancies_with_skills.csv')